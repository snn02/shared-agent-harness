"""Extract DOCX body blocks in document order without assuming page numbers."""

import hashlib
import io
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image

from ingest_model import ExtractedSource, Figure, Section


def extract_docx(path: Path) -> ExtractedSource:
    document = Document(path)
    headings: list[str] = []
    sections: list[Section] = []
    figures: list[Figure] = []
    warnings: list[str] = []
    seen_images: set[str] = set()
    lines: list[str] = []
    title = "Document"
    locator = "body"
    block_no = 0

    def flush() -> None:
        nonlocal lines
        if lines:
            sections.append(Section(f"section-{len(sections)+1:04}", title, locator,
                                    "\n\n".join(lines), "docx"))
            lines = []

    def scan_images(element, place: str) -> None:
        for blip in element.iter(qn("a:blip")):
            rel = blip.get(qn("r:embed"))
            if not rel or rel not in document.part.related_parts:
                warnings.append(f"Unresolved image at {place}")
                continue
            data = document.part.related_parts[rel].blob
            digest = hashlib.sha256(data).hexdigest()
            if digest in seen_images:
                continue
            seen_images.add(digest)
            try:
                buffer = io.BytesIO()
                with Image.open(io.BytesIO(data)) as image:
                    image.convert("RGB").save(buffer, format="PNG")
                figures.append(Figure(f"figure-{len(figures)+1:04}", place, buffer.getvalue()))
            except Exception:
                warnings.append(f"Unsupported image at {place}")

    def walk_table(table: Table, depth: int = 0) -> None:
        if depth > 10:
            warnings.append(f"Nested table exceeds depth 10 at block {block_no}")
            return
        for row in table.rows:
            cells = []
            for cell in row.cells:
                parts = []
                for child in cell._tc.iterchildren():
                    if child.tag == qn("w:p"):
                        value = Paragraph(child, cell).text.strip()
                        if value:
                            parts.append(value)
                    elif child.tag == qn("w:tbl"):
                        nested = Table(child, cell)
                        nested_rows = [" / ".join(c.text for c in r.cells) for r in nested.rows]
                        parts.append("; ".join(nested_rows))
                cells.append(" ".join(parts).replace("|", "\\|"))
            lines.append("| " + " | ".join(cells) + " |")

    for child in document.element.body.iterchildren():
        if child.tag not in (qn("w:p"), qn("w:tbl")):
            if child.tag != qn("w:sectPr"):
                warnings.append(f"Unsupported DOCX body element: {child.tag}")
            continue
        block_no += 1
        if child.tag == qn("w:tbl"):
            walk_table(Table(child, document))
            scan_images(child, f"{locator}, table block {block_no}")
            continue
        paragraph = Paragraph(child, document)
        style = paragraph.style.name or ""
        if style.startswith("Heading "):
            flush()
            try:
                level = int(style.split()[-1])
            except ValueError:
                level = 1
            headings = headings[:max(0, level - 1)] + [paragraph.text.strip()]
            title = paragraph.text.strip() or f"Heading at block {block_no}"
            locator = " / ".join(headings) + f" (block {block_no})"
        text = paragraph.text.strip()
        if text:
            lines.append(("#" * min(6, len(headings)) + " " if style.startswith("Heading ") else "") + text)
        scan_images(child, f"{locator}, block {block_no}")
        for marker in ("footnoteReference", "endnoteReference", "oMath", "object", "altChunk"):
            if any(e.tag == qn(f"w:{marker}") for e in child.iter()):
                warnings.append(f"Unsupported {marker} at block {block_no}")
    flush()
    for part in ("header", "footer"):
        if any(getattr(section, part).paragraphs for section in document.sections):
            # python-docx exposes an empty default paragraph even when no content exists.
            if any(p.text.strip() for section in document.sections for p in getattr(section, part).paragraphs):
                warnings.append(f"{part} text not extracted")
    if not sections:
        sections.append(Section("section-0001", title, locator, "", "docx"))
    return ExtractedSource(sections, figures, len(sections), warnings)
