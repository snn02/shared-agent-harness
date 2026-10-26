import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image

SCRIPTS = Path(__file__).resolve().parents[2] / "distribution/windows/payload/skills/document-ingest/scripts"
sys.path.insert(0, str(SCRIPTS))
from extract_docx import extract_docx

SCRATCH = Path(__file__).resolve().parents[2] / "dist" / "test-scratch"
SCRATCH.mkdir(parents=True, exist_ok=True)


class DocxTests(unittest.TestCase):
    def test_order_and_nested_table(self):
        with tempfile.TemporaryDirectory(dir=SCRATCH) as root:
            path = Path(root) / "sample.docx"
            doc = Document()
            doc.add_heading("Part", level=1)
            doc.add_paragraph("Before table")
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Cell text"
            doc.add_paragraph("After table")
            doc.save(path)
            result = extract_docx(path)
            content = "\n".join(s.markdown for s in result.sections)
            self.assertLess(content.index("Before table"), content.index("Cell text"))
            self.assertLess(content.index("Cell text"), content.index("After table"))
            self.assertIn("Part", result.sections[0].locator)

    def test_deduplicates_image_inside_table(self):
        with tempfile.TemporaryDirectory(dir=SCRATCH) as root:
            root = Path(root)
            image = root / "picture.png"
            Image.new("RGB", (10, 10), "red").save(image)
            path = root / "sample.docx"
            doc = Document()
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).paragraphs[0].add_run().add_picture(str(image))
            table.cell(0, 1).paragraphs[0].add_run().add_picture(str(image))
            doc.save(path)
            result = extract_docx(path)
            self.assertEqual(len(result.figures), 1)
            self.assertFalse(any("Unresolved image" in warning for warning in result.warnings))


if __name__ == "__main__":
    unittest.main()
